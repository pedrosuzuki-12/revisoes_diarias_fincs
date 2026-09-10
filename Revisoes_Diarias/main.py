import os
import re
import math
import uuid
import shutil
import logging
import requests
import pandas as pd
from datetime import datetime
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles

COLUNA_RETORNO = 'N'
TOLERANCIA = 0.0001

logging.basicConfig(level=logging.INFO, format='%(levelname)s: %(message)s')
logger = logging.getLogger(__name__)

app = FastAPI(title="Revisão de Pagamentos - Fincs")

# pasta temporária pra guardar uploads e resultados
TMP_DIR = "tmp_arquivos"
os.makedirs(TMP_DIR, exist_ok=True)


# ==========================================
# UTILITÁRIOS 
# ==========================================
def num_br(valor):
    if valor is None:
        return 0.0
    if isinstance(valor, (int, float)):
        return float(valor)
    try:
        texto = str(valor).replace(".", "").replace(",", ".").strip()
        return float(texto)
    except ValueError:
        return 0.0


# ==========================================
# OLIVEIRA TRUST 
# ==========================================
headers_ot = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36",
    "Referer": "https://www.oliveiratrust.com.br/",
    "Origin": "https://www.oliveiratrust.com.br",
    "Accept": "application/json, text/plain, */*",
}


def busca_id_oliveira(codigo_if):
    url = "https://services-ft.oliveiratrust.com.br/app/v1/titulos?busca=" + codigo_if
    try:
        resposta = requests.get(url, headers=headers_ot, timeout=10)
        resposta.raise_for_status()
        dados = resposta.json()
        lista_id = dados.get("data", [])
        if not lista_id:
            return None
        id = lista_id[0]
        return id.get("tit"), id.get("titulo")
    except Exception as e:
        logger.error(f"Erro na requisição OT: {e}")
        return None


def busca_historico_oliveira(id_ot, data_evento, max_paginas=100, data_limite="2000-01-01"):
    try:
        dt_alvo = datetime.strptime(data_evento, "%Y-%m-%d")
    except:
        return None
        
    melhor_item = None
    menor_distancia = 6

    for page in range(1, max_paginas + 1):
        url = f"https://services-ft.oliveiratrust.com.br/app/v1/titulos/historico_pu/{id_ot}?page={page}&limit=50"
        try:
            resposta = requests.get(url, headers=headers_ot, timeout=10)
            if resposta.status_code != 200:
                break
            dados = resposta.json()
            infos = dados.get("data", {}).get("data", [])
            if not infos:
                break

            for item in infos:
                if isinstance(item, dict):
                    # Filtra apenas dias com eventos reais
                    total_pgto = num_br(item.get("total_pgto"))
                    amort_pgto = num_br(item.get("amort_pgto"))
                    juros_pgto = num_br(item.get("juros_pgto"))
                    
                    if total_pgto > 0 or amort_pgto != 0 or juros_pgto > 0:
                        data_item_str = item.get("data")
                        if data_item_str:
                            try:
                                dt_item = datetime.strptime(data_item_str, "%Y-%m-%d")
                                dist = abs((dt_item - dt_alvo).days)
                                if dist < menor_distancia:
                                    menor_distancia = dist
                                    melhor_item = item
                                    if dist == 0:
                                        break
                            except:
                                pass
            
            if menor_distancia == 0:
                break

            ultima_data = infos[-1].get("data")
            if ultima_data:
                dt_ultima = datetime.strptime(ultima_data, "%Y-%m-%d")
                if (dt_alvo - dt_ultima).days > 5:
                    break
        except Exception:
            break

    if melhor_item:
        return (
            num_br(melhor_item.get("valor_nominal")), num_br(melhor_item.get("juros")),
            num_br(melhor_item.get("pu")), num_br(melhor_item.get("amort_pgto")),
            num_br(melhor_item.get("juros_pgto")), num_br(melhor_item.get("premio_pgto")),
            num_br(melhor_item.get("total_pgto"))
        )
    return None


def checagem_oliveira(amort_ord_per, amex_per, incorp_per, vn, juros, pu, amort_pgto, juros_pgto, premio_pgto, total_pgto, juros_per=100.0):
    esp_amort_ord = num_br(amort_ord_per)
    esp_amex = num_br(amex_per)
    esp_incorp = num_br(incorp_per)

    amort_real = (amort_pgto / vn) * 100.0 if vn > 0 and amort_pgto > 0 else 0.0
    amort_esperada = esp_amort_ord + esp_amex

    incorp_real = 0.0
    if amort_pgto < 0 and vn > 0:
        if abs(amort_pgto) == juros_pgto:
            incorp_real = 100.0
        else:
            incorp_real = (abs(amort_pgto) / juros_pgto) * 100
    elif total_pgto <= 0:
        incorp_real = 100.0

    if math.isclose(amort_real, amort_esperada, abs_tol=TOLERANCIA):
        amort_ord_real = esp_amort_ord
        amex_real = esp_amex
    else:
        if amort_real > esp_amort_ord:
            amort_ord_real = esp_amort_ord
            amex_real = amort_real - esp_amort_ord
        else:
            amort_ord_real = amort_real
            amex_real = 0.0

    status_geral = "OK"
    divergencias = []

    if not math.isclose(amort_real, amort_esperada, abs_tol=TOLERANCIA):
        status_geral = "DIVERGENTE"
        msg = f"Amort. Total Real ({amort_real:.8f}%) != Esperada ({amort_esperada:.8f}%)"
        if amort_real > amort_esperada:
            amex_calc = amort_real - amort_esperada
            msg += f" | AMEX detectada: {amex_calc:.8f}%"
        divergencias.append(msg)

    if not math.isclose(incorp_real, esp_incorp, abs_tol=TOLERANCIA):
        status_geral = "DIVERGENTE"
        str_incorp_real = f"{incorp_real:.8f}".replace(".00000000", "")
        str_esp_incorp = f"{esp_incorp:.8f}".replace(".00000000", "")
        divergencias.append(f"Incorp. Real ({str_incorp_real}%) != Esperada ({str_esp_incorp}%)")

    amort_ord_real_str = str(round(amort_ord_real, 4)).replace(".", ",")
    amex_real_str = str(round(amex_real, 4)).replace(".", ",")
    incorp_real_str = str(round(incorp_real, 4)).replace(".", ",")

    if status_geral == "OK":
        log_suscinto = f"OK (Amort. Ordinária: {amort_ord_real_str}% | AMEX: {amex_real_str}% | Incorporação: {incorp_real_str}%)"
    else:
        log_suscinto = "DIVERGÊNCIA - " + " | ".join(divergencias).replace(".", ",")

    return {"log_suscinto": log_suscinto}


# ==========================================
# VÓRTX 
# ==========================================
headers_vt = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
    "Accept": "application/json, text/plain, */*",
    "Origin": "https://www.vortx.com.br",
    "Referer": "https://www.vortx.com.br/",
}


def busca_id_vortx(codigo_cetip):
    codigo_cetip = str(codigo_cetip).strip()
    secoes = ["dcm", "fundos-de-investimento"]
    
    for secao in secoes:
        url = f"https://www.vortx.com.br/investidor/{secao}?busca={codigo_cetip}"
        try:
            response = requests.get(url, headers=headers_vt, timeout=15)
            if response.status_code == 200:
                soup = BeautifulSoup(response.text, "html.parser")
                link = soup.find("a", href=re.compile(r"/operacao\?id=\d+"))
                if link and "href" in link.attrs:
                    match = re.search(r"id=(\d+)", link["href"])
                    if match:
                        return match.group(1)
        except Exception:
            pass
            
    return None


def busca_historico_vortx(id_operacao, data_evento):
    url = f"https://apis.vortx.com.br/vxsite/api/operacao/{id_operacao}/preco-unitario/historico-pagamentos"
    try:
        response = requests.get(url, headers=headers_vt, timeout=15)
        response.raise_for_status()
        dados = response.json()
        if not dados.get("success") or "unitPrices" not in dados:
            return None

        df_pu = pd.DataFrame(dados["unitPrices"])
        if "paymentDate" in df_pu.columns:
            # Converte e remove o timezone para evitar erro de datas mistas (com e sem timezone)
            df_pu["paymentDate"] = pd.to_datetime(df_pu["paymentDate"], format="mixed", utc=True).dt.tz_localize(None)
            
            # Filtra apenas eventos reais (ignora a marcação diária de PU sem pagamento/evento)
            mask_evento = (df_pu["total"] > 0) | (df_pu["amortization"] != 0)
            if "unitPriceFull" in df_pu.columns and "unitPriceEmpty" in df_pu.columns:
                mask_evento = mask_evento | (df_pu["unitPriceFull"] != df_pu["unitPriceEmpty"])
                
            df_eventos = df_pu[mask_evento].copy()
            
            if not df_eventos.empty:
                dt_alvo = pd.to_datetime(data_evento)
                df_eventos["distancia"] = (df_eventos["paymentDate"] - dt_alvo).abs().dt.days
                
                df_filtrado = df_eventos[df_eventos["distancia"] <= 5]
                
                if not df_filtrado.empty:
                    melhor_indice = df_filtrado["distancia"].idxmin()
                    return df_eventos.loc[melhor_indice].copy()
                
    except Exception:
        pass
    return None


def checagem_vortx(amort_ord_per, amex_per, juros_per, incorp_per, linha):
    esp_amort_ord = num_br(amort_ord_per)
    esp_amex = num_br(amex_per)
    esp_incorp = num_br(incorp_per)

    if linha is None:
        return {"log_suscinto": "Sem evento na data escolhida ou Não pagou nada"}

    # Na Vortx, o PU Vazio vem na chave "unitPriceEmpty" e o valor nominal em "nominalValue"
    pu_vazio = float(linha.get("unitPriceEmpty", linha.get("nominalValue", 0))) 
    vlr_nominal = float(linha.get("nominalValue", 0))
    total_pago = float(linha.get("total", 0))
    amort_paga = float(linha.get("amortization", 0))
    juros_pago = float(linha.get("interestValue", 0.0))

    pu_base = pu_vazio + amort_paga
    amort_real = (amort_paga / pu_base) * 100.0 if pu_base > 0 and amort_paga > 0 else 0.0
    amort_esperada = esp_amort_ord + esp_amex

    incorp_real = 0.0
    if amort_paga < 0 and vlr_nominal > 0:
        if abs(amort_paga) == juros_pago:
            incorp_real = 100.0
        else:
            incorp_real = (abs(amort_paga) / juros_pago) * 100
    elif total_pago <= 0:
        incorp_real = 100.0

    if math.isclose(amort_real, amort_esperada, abs_tol=TOLERANCIA):
        amort_ord_real = esp_amort_ord
        amex_real = esp_amex
    else:
        if amort_real > esp_amort_ord:
            amort_ord_real = esp_amort_ord
            amex_real = amort_real - esp_amort_ord
        else:
            amort_ord_real = amort_real
            amex_real = 0.0

    status_geral = "OK"
    divergencias = []

    if not math.isclose(amort_real, amort_esperada, abs_tol=TOLERANCIA):
        status_geral = "DIVERGENTE"
        msg = f"Amort. Total Real ({amort_real:.8f}%) != Esperada ({amort_esperada:.8f}%)"
        if amort_real > amort_esperada:
            amex_calc = amort_real - amort_esperada
            msg += f" | AMEX detectada: {amex_calc:.8f}%"
        divergencias.append(msg)

    if not math.isclose(incorp_real, esp_incorp, abs_tol=TOLERANCIA):
        status_geral = "DIVERGENTE"
        str_incorp_real = f"{incorp_real:.8f}".replace(".00000000", "")
        str_esp_incorp = f"{esp_incorp:.8f}".replace(".00000000", "")
        divergencias.append(f"Incorp. Real ({str_incorp_real}%) != Esperada ({str_esp_incorp}%)")

    amort_ord_real_str = str(round(amort_ord_real, 4)).replace(".", ",")
    amex_real_str = str(round(amex_real, 4)).replace(".", ",")
    incorp_real_str = str(round(incorp_real, 4)).replace(".", ",")

    if status_geral == "OK":
        log_suscinto = f"OK (Amort. Ordinária: {amort_ord_real_str}% | AMEX: {amex_real_str}% | Incorporação: {incorp_real_str}%)"
    else:
        log_suscinto = "DIVERGÊNCIA - " + " | ".join(divergencias).replace(".", ",")

    return {"log_suscinto": log_suscinto}


# ==========================================
# PROCESSAMENTO 
# ==========================================
def processar_arquivo(caminho_entrada: str, caminho_saida: str):
    df = pd.read_excel(caminho_entrada, sheet_name=0, usecols="A:I")
    df = df.dropna(subset=[df.columns[0]])

    wb = load_workbook(caminho_entrada, keep_links=False)
    ws = wb.worksheets[0]

    for index, row in df.iterrows():
        linha_excel = index + 2
        codigo_if = str(row.iloc[0]).strip()
        data_evento = pd.to_datetime(row.iloc[2], dayfirst=True).strftime('%Y-%m-%d')

        amort_ord_per = num_br(row.iloc[3])
        amex_per = num_br(row.iloc[4])
        juros_per = num_br(row.iloc[6])
        incorp_per = num_br(row.iloc[7])
        af = str(row.iloc[8]).strip().upper()

        logger.info(f"Linha {linha_excel}: Título {codigo_if} ({af})")

        if af == "OLIVEIRA TRUST DTVM S.A.":
            res_id = busca_id_oliveira(codigo_if)
            if not res_id:
                ws[f"{COLUNA_RETORNO}{linha_excel}"] = "Título não encontrado na Oliveira Trust"
                continue

            id_ot, _ = res_id
            link_ot = f"https://www.oliveiratrust.com.br/investidor/ativos/historico-valores/{id_ot}"
            ws[f"M{linha_excel}"].value = link_ot
            ws[f"M{linha_excel}"].hyperlink = link_ot
            ws[f"M{linha_excel}"].style = "Hyperlink"

            dados_evento = busca_historico_oliveira(id_ot, data_evento)

            if dados_evento is None:
                ws[f"{COLUNA_RETORNO}{linha_excel}"] = "Nenhum evento encontrado na data"
            else:
                vn, juros, pu, amort_pgto, juros_pgto, premio_pgto, total_pgto = dados_evento
                resultado = checagem_oliveira(
                    amort_ord_per, amex_per, incorp_per, vn, juros, pu,
                    amort_pgto, juros_pgto, premio_pgto, total_pgto, juros_per
                )
                ws[f"{COLUNA_RETORNO}{linha_excel}"] = resultado["log_suscinto"]

        elif af == "VORTX DTVM LTDA.":
            id_vortx = busca_id_vortx(codigo_if)
            if not id_vortx:
                ws[f"{COLUNA_RETORNO}{linha_excel}"] = "Título não encontrado na Vórtx"
                continue

            link_vortx = f"https://www.vortx.com.br/investidor/dcm/operacao?id={id_vortx}"
            ws[f"M{linha_excel}"].value = link_vortx
            ws[f"M{linha_excel}"].hyperlink = link_vortx
            ws[f"M{linha_excel}"].style = "Hyperlink"

            dados_evento = busca_historico_vortx(id_vortx, data_evento)

            if dados_evento is None:
                ws[f"{COLUNA_RETORNO}{linha_excel}"] = "Nenhum evento encontrado na data"
            else:
                resultado = checagem_vortx(amort_ord_per, amex_per, juros_per, incorp_per, dados_evento)
                ws[f"{COLUNA_RETORNO}{linha_excel}"] = resultado["log_suscinto"]
        else:
            ws[f"{COLUNA_RETORNO}{linha_excel}"] = f"AF não mapeado: {af}"

    wb.save(caminho_saida)


# ==========================================
# ENDPOINTS
# ==========================================
@app.post("/processar")
async def processar(arquivo: UploadFile = File(...)):
    if not arquivo.filename.endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Envie um arquivo .xlsx")

    # cada requisição usa um id único, pra não misturar arquivos de gente diferente
    id_execucao = str(uuid.uuid4())
    caminho_entrada = os.path.join(TMP_DIR, f"{id_execucao}_input.xlsx")
    caminho_saida = os.path.join(TMP_DIR, f"{id_execucao}_output.xlsx")

    with open(caminho_entrada, "wb") as f:
        shutil.copyfileobj(arquivo.file, f)

    try:
        processar_arquivo(caminho_entrada, caminho_saida)
    except Exception as e:
        logger.error(f"Erro processando {arquivo.filename}: {e}")
        raise HTTPException(status_code=500, detail=f"Erro no processamento: {e}")
    finally:
        os.remove(caminho_entrada)

    nome_saida = arquivo.filename.replace(".xlsx", "_Retorno.xlsx")
    return FileResponse(
        path=caminho_saida,
        filename=nome_saida,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.get("/")
def home():
    return FileResponse("static/index.html")


app.mount("/static", StaticFiles(directory="static"), name="static")